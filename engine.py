from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass
from datetime import datetime
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# -----------------------------------------------------------------------------
# Rules carried from the supplied VBA package
# -----------------------------------------------------------------------------
LAND_ONLY_CODES = {
    "810", "811", "812", "813", "821", "822", "823", "831", "832", "833",
    "891", "910", "913", "914", "915", "916", "931",
}
STANDARD_CODES = {
    "111", "112", "113", "114", "116", "117", "118", "119", "120", "124",
    "130", "140", "141", "150", "155", "200", "300", "400", "460", "480",
    "500", "510", "511", "530", "531", "532", "533", "534", "535", "536",
    "537", "538", "539", "540", "550", "551", "553", "580", "583", "590",
    "600", "610", "611", "635", "636", "637", "638", "651", "652", "653",
    "654", "656", "657", "670", "671", "672", "680", "691", "721", "723",
    "730", "731", "740", "741", "748", "749", "750", "760", "850", "930",
    "951", "999",
}
MANUFACTURED_HOME_CODES = {"140", "141"}
APPROVED_CODES = LAND_ONLY_CODES | STANDARD_CODES | MANUFACTURED_HOME_CODES

BAD_SALE_KEYWORDS = [
    "MH", "MOBILE HOME", "TRAILER", "MANUFACTURED", "DOUBLEWIDE", "SINGLEWIDE",
    "MODULAR", "FORECLOSURE", "FORECLOSED", "REO", "HUD", "FANNIE MAE",
    "FREDDIE MAC", "SHORT SALE", "DISTRESSED", "AUCTION", "JUDICIAL", "SHERIFF",
    "LIEN", "DEED IN LIEU", "BANKRUPTCY", "ESTATE", "TRUST", "TRUSTEE",
    "EXECUTOR", "EXECUTRIX", "HEIRS", "DECEASED", "PERSONAL REPRESENTATIVE",
    "INTRA-FAMILY", "GIFT", "RELATED PARTIES", "PARENT TO CHILD", "SIBLINGS",
    "CORPORATE TRANSFER", "AFFILIATE", "SUBSIDIARY", "PARTNERSHIP", "HOLDINGS",
    "INC", "CORP", "TAX DEED", "TAX SALE", "GOVERNMENT", "HOUSING AUTHORITY",
    "DEVELOPMENT AUTHORITY", "CITY OF", "COUNTY OF", "STATE OF", "BOARD OF EDUCATION",
    "CHURCH", "PRAYER", "PRAYED", "MINISTRIES", "NON-PROFIT", "PRAY",
    "EMINENT DOMAIN", "CONDEMNATION", "EASEMENT", "RIGHT OF WAY", "CORRECTION DEED",
    "QUIT CLAIM", "QUITCLAIM", "LOVE AND AFFECTION", "NOMINAL", "PARTIAL INTEREST",
    "UNDIVIDED INTEREST", "LIFE ESTATE", "GUARDIAN", "CONSERVATOR", "LEASEBACK",
    "TRADE", "BARTER", "EXCHANGE", "DESTROYED", "BURNED", "FIRE", "DEMOLISHED",
    "UNINHABITABLE", "CONDEMNED", "BLIGHT", "FLOODED", "SPLIT", "COMBINATION",
    "REVISED PARCEL", "UNDER CONSTRUCTION", "R/E", "L/E",
]

GLOBAL_TOO_LOW = 0.50
GLOBAL_PERFECT_LOW = 0.70
GLOBAL_PERFECT_HIGH = 1.20
GLOBAL_TOO_HIGH = 1.50
CROSSFOOT_TOLERANCE = 0.01
MIN_REQUIRED_SALES = 10
NATURAL_MIN_RATIO = 0.975
NATURAL_MAX_RATIO = 1.0244
TREND_MIN_RATIO = 0.995
TREND_MAX_RATIO = 1.0044
MAX_COD = 0.20

GOOD_VALUES = {
    "Y", "YES", "TRUE", "1", "GOOD", "GOODSALE", "INCLUDE", "INCLUDED",
    "INCLUDEDINSTUDY", "QUALIFIED", "USABLE", "USEABLE", "VALID", "RATIOABLE",
}
BAD_VALUES = {
    "N", "NO", "FALSE", "0", "BAD", "BADSALE", "EXCLUDE", "EXCLUDED", "REJECT",
    "REJECTED", "UNQUALIFIED", "NOTUSABLE", "INVALID",
}

HEADER_ALIASES = {
    "parcel": ["PARCEL", "PARCELS", "PARCELNUMBER", "PARCELNO", "PCL", "PCLNUMBER"],
    "sale_price": ["SALEPRICE", "SALES PRICE", "SALESPRICE", "PRICE"],
    "land_value": ["LANDVALUE", "LANDVAL"],
    "improvement_value": ["IMPROVEMENTVALUE", "BLDGVALUE", "BLDVALUE", "BUILDINGVALUE"],
    "misc_value": ["MISCIMPROVEMENTVALUE", "MISCELLANEOUSIMPROVEMENTVALUE", "MISCIMPROVEMENTSVALUE", "MISCIMPROVEMENTS", "MISCVALUE", "MISCELLANEOUSVALUE"],
    "total_value": ["TOTALVALUE", "TOTALVALUEOFPROPERTY", "TOTALAPPRAISAL", "APPRAISEDVALUE", "TOTALPROPERTYVALUE"],
    "use_code": ["USECODE", "LANDUSECODE", "LANDUSE", "LANDUSECODES", "USESUBCODE", "IMPROVEMENTCODE"],
    "type": ["TYPE", "SALETYPE"],
    "ratio": ["RATIO", "SALESRATIO", "SALERATIO"],
    "qualification": ["QUALIFIEDSALE", "INCLUDEDSALE", "QUALIFICATIONSTATUS", "QUALIFICATION", "QUALIFIED", "SALEQUALIFICATION"],
    "neighborhood": ["VALUATIONZONE", "NEIGHBORHOOD", "NBHD", "NBRHD", "NBHCODE", "ZONECODE"],
    "grantor": ["GRANTOR", "GONAME", "GRANTORNAME"],
    "grantee": ["GRANTEE", "GENAME", "GRANTEENAME"],
    "comment": ["SALENOTES", "COMMENTS", "COMMENT", "SALECOMMENT1", "SALECOMMENT2", "SALECOMMENT3", "NOTES"],
    "sale_date": ["DDATE", "SALEDATE"],
    "sale_date_year": ["SALEDATEYEAR"],
    "sale_date_month": ["SALEDATEMONTH"],
    "sale_date_day": ["SALEDATEDAY"],
    "deed_book": ["DEEDBOOK", "DBOOK", "RECDBOOK", "BOOK"],
    "deed_page": ["DEEDPAGE", "DEEDBOOKPAGE", "DBPAGE", "DEPAGE", "RECDPAGE", "PAGE"],
}

STATUS_COL = "Flag Status"
NEIGHBORHOOD_CHECK_COL = "Neighborhood Check"
RATIO_COL = "Sales Ratio"
TYPE_COL = "Type"


def normalize_header(value: object) -> str:
    s = "" if value is None else str(value).strip().upper()
    for ch in (" ", "_", "-", "/", "."):
        s = s.replace(ch, "")
    return s


def normalize_code(value: object) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    s = str(value).strip()
    # Excel often imports numeric codes as 111.0.
    try:
        f = float(s)
        if math.isfinite(f) and f.is_integer():
            return str(int(f))
    except Exception:
        pass
    return s.upper().replace(" ", "").replace("-", "").replace("_", "")


def clean_text(value: object) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    return str(value).strip()


def numeric(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce").fillna(0.0)


def _colmap(df: pd.DataFrame) -> Dict[str, List[str]]:
    out: Dict[str, List[str]] = {}
    for c in df.columns:
        out.setdefault(normalize_header(c), []).append(c)
    return out


def columns_for(df: pd.DataFrame, logical_name: str) -> List[str]:
    norm_map = _colmap(df)
    found: List[str] = []
    for alias in HEADER_ALIASES.get(logical_name, []):
        found.extend(norm_map.get(normalize_header(alias), []))
    # Preserve original column order and uniqueness.
    return [c for c in df.columns if c in set(found)]


def first_col(df: pd.DataFrame, logical_name: str) -> Optional[str]:
    cols = columns_for(df, logical_name)
    return cols[0] if cols else None


def sum_columns(df: pd.DataFrame, cols: Sequence[str]) -> pd.Series:
    if not cols:
        return pd.Series(0.0, index=df.index)
    return pd.concat([numeric(df[c]) for c in cols], axis=1).sum(axis=1)


def classify_qualification_value(value: object) -> str:
    raw = clean_text(value)
    if not raw:
        return "UNKNOWN"
    norm = normalize_header(raw)
    if norm in GOOD_VALUES or "GOODSALE" in norm or "INCLUDEDINSTUDY" in norm:
        return "GOOD"
    if norm in BAD_VALUES or "BADSALE" in norm or "REJECT" in norm or "EXCLUDE" in norm:
        return "BAD"
    return "UNKNOWN"


def classify_row_qualification(df: pd.DataFrame, idx) -> Tuple[str, str]:
    cols = columns_for(df, "qualification")
    if not cols:
        return "UNKNOWN", "No qualification field mapped"
    statuses = []
    details = []
    for c in cols:
        raw = clean_text(df.at[idx, c])
        status = classify_qualification_value(raw)
        statuses.append(status)
        if raw:
            details.append(f"{c}={raw}")
    has_good = "GOOD" in statuses
    has_bad = "BAD" in statuses
    if has_good and has_bad:
        return "CONFLICT", "; ".join(details)
    if has_good:
        return "GOOD", "; ".join(details)
    if has_bad:
        return "BAD", "; ".join(details)
    return "UNKNOWN", "; ".join(details)


def read_uploaded_workbook(file_bytes: bytes, filename: str) -> Tuple[pd.DataFrame, str]:
    """Read the Analysis sheet if present; otherwise the first sheet."""
    lower = filename.lower()
    bio = io.BytesIO(file_bytes)
    if lower.endswith(".xlsb"):
        from pyxlsb import open_workbook
        # pyxlsb needs a path or file-like object depending on environment. BytesIO works in 1.0.x.
        with open_workbook(bio) as wb:
            sheet_name = "Analysis" if "Analysis" in wb.sheets else wb.sheets[0]
            with wb.get_sheet(sheet_name) as sh:
                rows = [[cell.v for cell in row] for row in sh.rows()]
        if not rows:
            return pd.DataFrame(), sheet_name
        width = max(len(r) for r in rows)
        rows = [r + [None] * (width - len(r)) for r in rows]
        headers = [clean_text(v) or f"Unnamed_{i+1}" for i, v in enumerate(rows[0])]
        return pd.DataFrame(rows[1:], columns=headers), sheet_name
    if lower.endswith(".xls"):
        xls = pd.ExcelFile(bio, engine="xlrd")
    else:
        xls = pd.ExcelFile(bio, engine="openpyxl")
    sheet_name = "Analysis" if "Analysis" in xls.sheet_names else xls.sheet_names[0]
    return pd.read_excel(xls, sheet_name=sheet_name, dtype=object), sheet_name


def prepare_data(df: pd.DataFrame) -> Tuple[pd.DataFrame, Dict[str, object]]:
    out = df.copy()
    before_rows, before_cols = out.shape
    out = out.dropna(how="all").copy()
    # Drop completely empty columns and unnamed placeholders with no values.
    out = out.dropna(axis=1, how="all")
    # Clean duplicate generated columns, retaining the first meaningful one later.
    generated_norms = {
        "FLAGSTATUS", "NEIGHBORHOODCHECK", "TYPEDEFAULT", "TYPEOVERRIDEFLAG", "SALETYPE",
        "VIEWDEFENSE", "DEFENSEDRILLHIGHLIGHT", "STANDARDSALEDATE", "STANDARDIZEDSALEDATE",
    }
    keep_cols = []
    seen_gen = set()
    for c in out.columns:
        n = normalize_header(c)
        if n in generated_norms:
            if n in seen_gen:
                continue
            seen_gen.add(n)
        keep_cols.append(c)
    out = out.loc[:, keep_cols].copy()

    # Consolidate split sale date pieces into one Sale Date when all three exist.
    y, m, d = first_col(out, "sale_date_year"), first_col(out, "sale_date_month"), first_col(out, "sale_date_day")
    if y and m and d:
        yy = pd.to_numeric(out[y], errors="coerce")
        mm = pd.to_numeric(out[m], errors="coerce")
        dd = pd.to_numeric(out[d], errors="coerce")
        combined = pd.to_datetime(dict(year=yy, month=mm, day=dd), errors="coerce")
        existing = first_col(out, "sale_date")
        if existing:
            out[existing] = combined
        else:
            out["Sale Date"] = combined

    # If there are many CODE columns, mimic the workbook's preferred-code consolidation.
    code_cols = [c for c in out.columns if "CODE" in normalize_header(c)]
    if len(code_cols) > 3:
        def valid_count(c):
            vals = out[c].map(normalize_code)
            return int(vals.isin(APPROVED_CODES).sum()), int((vals != "").sum())
        ranked = sorted(code_cols, key=lambda c: valid_count(c), reverse=True)
        preferred = ranked[0]
        values = []
        for idx in out.index:
            selected = clean_text(out.at[idx, preferred])
            if normalize_code(selected) not in APPROVED_CODES:
                for c in code_cols:
                    raw = clean_text(out.at[idx, c])
                    if normalize_code(raw) in APPROVED_CODES:
                        selected = normalize_code(raw)
                        break
                    if not selected and raw:
                        selected = raw
            values.append(selected)
        out["Use Code"] = values

    # Preserve identifier-like columns as strings.
    for c in out.columns:
        n = normalize_header(c)
        if "PARCEL" in n or n.startswith("PCL"):
            out[c] = out[c].map(lambda v: "" if pd.isna(v) else str(v))

    # Remove prior diagnostic statuses so the next checker starts clean.
    for c in [STATUS_COL, NEIGHBORHOOD_CHECK_COL]:
        if c in out.columns:
            out.drop(columns=[c], inplace=True)

    report = {
        "rows_before": before_rows,
        "rows_after": len(out),
        "columns_before": before_cols,
        "columns_after": out.shape[1],
        "blank_rows_removed": before_rows - len(out),
    }
    return out.reset_index(drop=True), report


def sale_date_checker(df: pd.DataFrame, tax_year: int) -> Tuple[pd.DataFrame, Dict[str, object]]:
    out = df.copy()
    date_col = first_col(out, "sale_date")
    y, m, d = first_col(out, "sale_date_year"), first_col(out, "sale_date_month"), first_col(out, "sale_date_day")
    mapped = date_col
    if not date_col and y and m and d:
        yy = pd.to_numeric(out[y], errors="coerce")
        mm = pd.to_numeric(out[m], errors="coerce")
        dd = pd.to_numeric(out[d], errors="coerce")
        out["Sale Date"] = pd.to_datetime(dict(year=yy, month=mm, day=dd), errors="coerce")
        date_col = "Sale Date"
        mapped = "SALE_DATE_YEAR / SALE_DATE_MONTH / SALE_DATE_DAY (Combined)"
    if not date_col:
        raise ValueError("No supported sale-date structure found. Expected D_DATE, Sale Date, or SALE_DATE_YEAR/MONTH/DAY.")

    raw = out[date_col]
    parsed = pd.to_datetime(raw, errors="coerce")
    start = pd.Timestamp(year=tax_year - 2, month=10, day=1)
    end = pd.Timestamp(year=tax_year - 1, month=9, day=30)
    blank = raw.isna() | raw.astype(str).str.strip().isin(["", "None", "nan", "NaT"])
    valid = parsed.notna() & (parsed >= start) & (parsed <= end)
    statuses = np.where(blank, "missing sales date information", np.where(valid, "Verified Active Sale", "sale date invalid"))
    out[STATUS_COL] = statuses
    summary = {
        "checker": "Sale Date Checker",
        "tax_year": tax_year,
        "study_start": start.date().isoformat(),
        "study_end": end.date().isoformat(),
        "mapped_field": mapped,
        "total_sales": len(out),
        "valid_sales": int((out[STATUS_COL] == "Verified Active Sale").sum()),
        "invalid_sales": int((out[STATUS_COL] == "sale date invalid").sum()),
        "missing_sales": int((out[STATUS_COL] == "missing sales date information").sum()),
    }
    return out, summary


def use_code_checker(df: pd.DataFrame) -> Tuple[pd.DataFrame, Dict[str, object]]:
    out = df.copy()
    use_cols = columns_for(out, "use_code")
    land_cols = columns_for(out, "land_value")
    imp_cols = columns_for(out, "improvement_value")
    misc_cols = columns_for(out, "misc_value")
    if not use_cols:
        raise ValueError("No supported Use Code field was found.")
    if not land_cols:
        raise ValueError("No Land Value field was found.")

    land_total = sum_columns(out, land_cols)
    imp_total = sum_columns(out, imp_cols)
    misc_total = sum_columns(out, misc_cols)
    statuses = []
    counters = dict(compliant=0, blank=0, invalid=0, valuation=0, mh=0)
    for pos, idx in enumerate(out.index):
        raw_codes = [clean_text(out.at[idx, c]) for c in use_cols]
        codes = [normalize_code(v) for v in raw_codes if v]
        issues = []
        if not codes:
            issues.append("Review: Missing Use Code Entry")
            counters["blank"] += 1
        else:
            invalid = [raw for raw, c in zip(raw_codes, [normalize_code(v) for v in raw_codes]) if raw and c not in APPROVED_CODES]
            if invalid:
                issues.append("Review Invalid Use Code: " + "; ".join(invalid))
                counters["invalid"] += 1
        land_only = any(c in LAND_ONLY_CODES for c in codes)
        mh = any(c in MANUFACTURED_HOME_CODES for c in codes)
        if float(land_total.iloc[pos]) <= 0:
            issues.append("Error: Missing or Zero Land Value")
        if land_only and float(imp_total.iloc[pos]) > 0:
            issues.append("Error: Improvement Value Found on Land-Only Code")
        if mh:
            issues.append("Review for Manufactured Home")
            counters["mh"] += 1
        if float(misc_total.iloc[pos]) < 0:
            issues.append("Review: Negative Miscellaneous Improvement Value")
        if any(x.startswith("Error:") for x in issues):
            counters["valuation"] += 1
        if not issues:
            counters["compliant"] += 1
            statuses.append("Verified Compliant")
        else:
            statuses.append("; ".join(issues))
    out[STATUS_COL] = statuses
    summary = {
        "checker": "Use Code Checker",
        "total_rows": len(out),
        "verified_compliant": counters["compliant"],
        "valuation_conflict_errors": counters["valuation"],
        "manufactured_home_reviews": counters["mh"],
        "unrecognized_use_codes": counters["invalid"],
        "missing_use_codes": counters["blank"],
        "issues_requiring_review": counters["valuation"] + counters["mh"] + counters["invalid"] + counters["blank"],
        "mapped_use_code_fields": ", ".join(use_cols),
    }
    return out, summary


def appraisal_value_checker(df: pd.DataFrame) -> Tuple[pd.DataFrame, Dict[str, object]]:
    out = df.copy()
    land_cols = columns_for(out, "land_value")
    imp_cols = columns_for(out, "improvement_value")
    misc_cols = columns_for(out, "misc_value")
    total_cols = columns_for(out, "total_value")
    if not land_cols or not total_cols:
        raise ValueError("Appraisal Value Checker requires Land Value and Total Value fields.")
    land = sum_columns(out, land_cols).round(2)
    imp = sum_columns(out, imp_cols).round(2)
    misc = sum_columns(out, misc_cols).round(2)
    stated = sum_columns(out, total_cols).round(2)
    calc = (land + imp + misc).round(2)
    variance = (calc - stated).round(2)
    statuses = []
    missing = mismatch = compliant = 0
    for i in range(len(out)):
        if land.iloc[i] == 0 or stated.iloc[i] == 0:
            missing += 1
            statuses.append("Error: Missing or Zero Land Value / Total Value")
        elif abs(float(variance.iloc[i])) > CROSSFOOT_TOLERANCE:
            mismatch += 1
            statuses.append(f"Error: Appraisal Value Mismatch (Calculated {calc.iloc[i]:,.2f}; Stated {stated.iloc[i]:,.2f}; Variance {variance.iloc[i]:,.2f})")
        else:
            compliant += 1
            statuses.append("Verified Compliant")
    out[STATUS_COL] = statuses
    out["Calculated Appraisal Total"] = calc
    out["Appraisal Variance"] = variance
    summary = {
        "checker": "Appraisal Value Checker",
        "total_rows": len(out),
        "verified_compliant": compliant,
        "missing_core_values": missing,
        "crossfoot_mismatches": mismatch,
        "issues_requiring_review": missing + mismatch,
        "tolerance": CROSSFOOT_TOLERANCE,
    }
    return out, summary


def _contains_keyword(text: str) -> Optional[str]:
    upper = text.upper()
    for kw in BAD_SALE_KEYWORDS:
        if re.search(r"(?<![A-Z0-9])" + re.escape(kw) + r"(?![A-Z0-9])", upper, flags=re.I):
            return kw
    return None


def deed_audit_checker(df: pd.DataFrame) -> Tuple[pd.DataFrame, Dict[str, object]]:
    out = df.copy()
    grantor_cols, grantee_cols = columns_for(out, "grantor"), columns_for(out, "grantee")
    comment_cols = columns_for(out, "comment")
    book_cols, page_cols = columns_for(out, "deed_book"), columns_for(out, "deed_page")
    nbhd_cols = columns_for(out, "neighborhood")
    date_col = first_col(out, "sale_date")
    if not grantor_cols or not grantee_cols:
        raise ValueError("Deed Audit requires Grantor and Grantee fields.")

    def joined(idx, cols):
        return " | ".join([clean_text(out.at[idx, c]) for c in cols if clean_text(out.at[idx, c])])

    statuses = [""] * len(out)
    counts: Dict[str, int] = {}
    transaction_first: Dict[Tuple[str, str, str], int] = {}
    deed_first: Dict[Tuple[str, str], Tuple[str, int]] = {}

    def bump(k): counts[k] = counts.get(k, 0) + 1

    for p, idx in enumerate(out.index):
        grantor = joined(idx, grantor_cols)
        grantee = joined(idx, grantee_cols)
        comment = joined(idx, comment_cols)
        book = joined(idx, book_cols)
        page = joined(idx, page_cols)
        nbhd = joined(idx, nbhd_cols)
        date_txt = clean_text(out.at[idx, date_col]) if date_col else ""
        qual, detail = classify_row_qualification(out, idx)

        # Priority 1: missing party.
        if not grantor or not grantee:
            statuses[p] = "Critical: Missing Grantor or Grantee"
            bump("missing_party"); continue
        # Priority 2: bad sale no comment.
        if qual == "BAD" and not comment:
            statuses[p] = "Critical: Bad Sale with No Comment"
            bump("bad_no_comment"); continue
        # Priority 3: qualification conflict.
        if qual == "CONFLICT":
            statuses[p] = f"Review: Conflicting Qualification Status [{detail}]"
            bump("qualification_conflict"); continue
        # Priority 4: unknown qualification.
        if qual == "UNKNOWN":
            statuses[p] = "Unknown Sale Status"
            bump("unknown_qualification"); continue
        # Priority 5: missing deed or neighborhood.
        if not book or not page or not nbhd:
            statuses[p] = "Review: Missing Deed / Neighborhood Information"
            bump("missing_reference"); continue
        # Priority 6: same deed in different neighborhood.
        dkey = (book.upper(), page.upper())
        if dkey in deed_first and deed_first[dkey][0].upper() != nbhd.upper():
            statuses[p] = f"Review: Neighborhood Mismatch [vs Row {deed_first[dkey][1] + 2}]"
            bump("neighborhood_mismatch"); continue
        deed_first.setdefault(dkey, (nbhd, p))
        # Priority 7: bad sale with comment.
        if qual == "BAD":
            statuses[p] = "Bad Sale with Comment - Review"
            bump("bad_with_comment"); continue
        # Priority 8: duplicate grantor/grantee/date.
        tkey = (grantor.upper(), grantee.upper(), date_txt.upper())
        if all(tkey) and tkey in transaction_first:
            statuses[p] = "Review: Possible Duplicate Transaction on Same Date"
            bump("duplicate_transaction"); continue
        if all(tkey): transaction_first[tkey] = p
        # Priority 9: self transfer.
        if grantor.upper() == grantee.upper():
            statuses[p] = "Review: Deed Correction / Self-Transfer"
            bump("self_transfer"); continue
        # Priority 10: manufactured home.
        if re.search(r"\b(MH|MOBILE HOME|TRAILER|MANUFACTURED|DOUBLEWIDE|SINGLEWIDE|MODULAR)\b", comment, re.I):
            statuses[p] = "Manufactured Home Review"
            bump("manufactured_home"); continue
        # Priority 11: LLC-to-LLC.
        if re.search(r"\bLLC\b", grantor, re.I) and re.search(r"\bLLC\b", grantee, re.I):
            statuses[p] = "Review: LLC-to-LLC Transaction"
            bump("llc_transfer"); continue
        # Priority 12: configured keyword.
        field_hit = None
        for label, txt in (("Grantor", grantor), ("Grantee", grantee), ("Comment", comment)):
            kw = _contains_keyword(txt)
            if kw:
                field_hit = f"Review: {label} Keyword [{kw}]"
                break
        if field_hit:
            statuses[p] = field_hit
            bump("keyword_review"); continue
        statuses[p] = "Verified Compliant"
        bump("compliant")

    out[STATUS_COL] = statuses
    summary = {"checker": "Deed Audit, MH, & Bad Sale Comment Checker", "total_rows": len(out), **counts}
    summary["issues_requiring_review"] = len(out) - counts.get("compliant", 0)
    return out, summary


def _default_type(use_code: str) -> str:
    if use_code in LAND_ONLY_CODES:
        return "L"
    if use_code in APPROVED_CODES:
        return "L&B"
    return ""


def ratio_checker(df: pd.DataFrame) -> Tuple[pd.DataFrame, Dict[str, object]]:
    out = df.copy()
    price_col = first_col(out, "sale_price")
    land_cols, imp_cols = columns_for(out, "land_value"), columns_for(out, "improvement_value")
    misc_cols, total_cols = columns_for(out, "misc_value"), columns_for(out, "total_value")
    use_col, type_col, nbhd_col = first_col(out, "use_code"), first_col(out, "type"), first_col(out, "neighborhood")
    if not price_col or not land_cols or not total_cols:
        raise ValueError("Ratio Checker requires Sale Price, Land Value, and Total Value.")
    price, land, imp, misc, total = numeric(out[price_col]), sum_columns(out, land_cols), sum_columns(out, imp_cols), sum_columns(out, misc_cols), sum_columns(out, total_cols)
    types, ratios, qstatuses = [], [], []
    for p, idx in enumerate(out.index):
        qual, _ = classify_row_qualification(out, idx)
        qstatuses.append(qual)
        existing = clean_text(out.at[idx, type_col]).upper() if type_col else ""
        if existing not in {"L", "B", "L&B"}:
            existing = _default_type(normalize_code(out.at[idx, use_col])) if use_col else ""
        types.append(existing)
        r = np.nan
        if qual == "GOOD" and price.iloc[p] > 0:
            if existing == "L": r = land.iloc[p] / price.iloc[p]
            elif existing == "B": r = (imp.iloc[p] + misc.iloc[p]) / price.iloc[p]
            elif existing == "L&B": r = total.iloc[p] / price.iloc[p]
        ratios.append(r)
    out[TYPE_COL] = types
    out[RATIO_COL] = ratios

    # Neighborhood quartiles based on valid good-sale ratios.
    q1q3: Dict[str, Tuple[float, float]] = {}
    if nbhd_col:
        temp = pd.DataFrame({"nbhd": out[nbhd_col].map(clean_text), "ratio": ratios, "qual": qstatuses})
        for nbhd, grp in temp[(temp["qual"] == "GOOD") & pd.to_numeric(temp["ratio"], errors="coerce").notna()].groupby("nbhd"):
            vals = pd.to_numeric(grp["ratio"], errors="coerce").dropna()
            if len(vals): q1q3[nbhd] = (float(vals.quantile(.25)), float(vals.quantile(.75)))

    flags, neighborhood_checks = [], []
    for p, idx in enumerate(out.index):
        r = ratios[p]
        qual = qstatuses[p]
        if qual != "GOOD" or not np.isfinite(r):
            flags.append("N/A - Sale Not Included / Ratio Unavailable")
            neighborhood_checks.append("N/A")
            continue
        if r < GLOBAL_TOO_LOW: global_status = "Too Low for Global"
        elif GLOBAL_PERFECT_LOW <= r <= GLOBAL_PERFECT_HIGH: global_status = "Perfect Global"
        elif r > GLOBAL_TOO_HIGH: global_status = "Too High for Global"
        else: global_status = "Acceptable Global"
        nstatus = "N/A"
        if nbhd_col:
            nbhd = clean_text(out.at[idx, nbhd_col])
            if nbhd in q1q3:
                q1, q3 = q1q3[nbhd]
                if r < q1: nstatus = "Too Low"
                elif r > q3: nstatus = "Too High"
                else: nstatus = "Acceptable"
        neighborhood_checks.append(nstatus)
        flags.append(f"{global_status}; Neighborhood: {nstatus}")
    out[STATUS_COL] = flags
    out[NEIGHBORHOOD_CHECK_COL] = neighborhood_checks
    valid = pd.to_numeric(out[RATIO_COL], errors="coerce").dropna()
    summary = {
        "checker": "Ratio Checker",
        "total_rows": len(out),
        "valid_ratios": int(valid.count()),
        "median_ratio": float(valid.median()) if len(valid) else 0.0,
        "mean_ratio": float(valid.mean()) if len(valid) else 0.0,
        "too_low_global": int(out[STATUS_COL].astype(str).str.contains("Too Low for Global", regex=False).sum()),
        "too_high_global": int(out[STATUS_COL].astype(str).str.contains("Too High for Global", regex=False).sum()),
        "perfect_global": int(out[STATUS_COL].astype(str).str.contains("Perfect Global", regex=False).sum()),
    }
    return out, summary


def quality_checker(df: pd.DataFrame) -> Tuple[pd.DataFrame, Dict[str, object]]:
    # Start from type-aware ratio results, then apply the broader integrity audit hierarchy.
    ratio_df, ratio_summary = ratio_checker(df)
    out = ratio_df.copy()
    parcel_col, nbhd_col, use_col = first_col(out, "parcel"), first_col(out, "neighborhood"), first_col(out, "use_code")
    price_col = first_col(out, "sale_price")
    date_col, grantor_col, grantee_col = first_col(out, "sale_date"), first_col(out, "grantor"), first_col(out, "grantee")
    comment_cols = columns_for(out, "comment")
    book_col, page_col = first_col(out, "deed_book"), first_col(out, "deed_page")
    land_cols, imp_cols, misc_cols, total_cols = columns_for(out, "land_value"), columns_for(out, "improvement_value"), columns_for(out, "misc_value"), columns_for(out, "total_value")
    land, imp, misc, total = sum_columns(out, land_cols), sum_columns(out, imp_cols), sum_columns(out, misc_cols), sum_columns(out, total_cols)

    # Calculate ratio quartiles for dual-outlier logic.
    valid_ratios = pd.to_numeric(out[RATIO_COL], errors="coerce")
    q1q3 = {}
    if nbhd_col:
        temp = pd.DataFrame({"nbhd": out[nbhd_col].map(clean_text), "ratio": valid_ratios})
        for nbhd, grp in temp.groupby("nbhd"):
            vals = grp["ratio"].dropna()
            if len(vals): q1q3[nbhd] = (float(vals.quantile(.25)), float(vals.quantile(.75)))

    statuses = []
    counts: Dict[str, int] = {}
    signatures, deeds = {}, {}
    def bump(k): counts[k] = counts.get(k, 0) + 1

    for p, idx in enumerate(out.index):
        qual, qdetail = classify_row_qualification(out, idx)
        parcel = clean_text(out.at[idx, parcel_col]) if parcel_col else ""
        nbhd = clean_text(out.at[idx, nbhd_col]) if nbhd_col else ""
        use = normalize_code(out.at[idx, use_col]) if use_col else ""
        grantor = clean_text(out.at[idx, grantor_col]) if grantor_col else ""
        grantee = clean_text(out.at[idx, grantee_col]) if grantee_col else ""
        date = clean_text(out.at[idx, date_col]) if date_col else ""
        book = clean_text(out.at[idx, book_col]) if book_col else ""
        page = clean_text(out.at[idx, page_col]) if page_col else ""
        comment = " | ".join(clean_text(out.at[idx, c]) for c in comment_cols if clean_text(out.at[idx, c]))
        ratio = valid_ratios.iloc[p]

        # Structural priority first (matching source description).
        if qual == "BAD" and not comment:
            status = "Critical: Bad Sale with No Comment"; bump("bad_sale_no_comment")
        elif qual == "CONFLICT":
            status = f"Review: Conflicting Qualification Status [{qdetail}]"; bump("qualification_conflict")
        elif qual == "UNKNOWN":
            status = "Unknown Sale Status"; bump("unknown_qualification")
        elif not parcel:
            status = "Critical: Missing Parcel Number"; bump("missing_parcel")
        elif not nbhd:
            status = "Critical: Missing Neighborhood / Valuation Zone"; bump("missing_neighborhood")
        elif not date:
            status = "Critical: Missing Sale Date"; bump("missing_sale_date")
        elif not use:
            status = "Critical: Missing Use Code"; bump("missing_use_code")
        elif use not in APPROVED_CODES:
            status = f"Review: Invalid / Unsupported Use Code [{use}]"; bump("invalid_use_code")
        elif land.iloc[p] <= 0:
            status = "Critical: Missing Land Value"; bump("missing_land_value")
        elif abs(float((land.iloc[p] + imp.iloc[p] + misc.iloc[p]) - total.iloc[p])) > CROSSFOOT_TOLERANCE:
            status = "Review: Appraisal Component-to-Total Value Imbalance"; bump("appraisal_imbalance")
        elif not grantor or not grantee:
            status = "Critical: Missing Grantor or Grantee"; bump("missing_party")
        elif book and page:
            dkey = (book.upper(), page.upper())
            if dkey in deeds and deeds[dkey][0].upper() != nbhd.upper():
                status = "Review: Same Deed Appears in Different Neighborhoods"; bump("deed_neighborhood_mismatch")
            else:
                deeds.setdefault(dkey, (nbhd, p))
                status = ""
        else:
            status = ""

        if not status:
            sig = (grantor.upper(), grantee.upper(), date.upper())
            if all(sig) and sig in signatures:
                status = "Review: Possible Duplicate Transaction"; bump("duplicate_transaction")
            elif all(sig):
                signatures[sig] = p

        if not status and qual == "BAD":
            status = "Bad Sale with Comment - Review"; bump("bad_sale_with_comment")
        if not status and np.isfinite(ratio):
            global_outlier = ratio < GLOBAL_TOO_LOW or ratio > GLOBAL_TOO_HIGH
            neighborhood_outlier = False
            if nbhd in q1q3:
                q1, q3 = q1q3[nbhd]
                neighborhood_outlier = ratio < q1 or ratio > q3
            if global_outlier and neighborhood_outlier:
                status = "Critical: Global and Neighborhood Ratio Outlier"; bump("dual_outlier")
            elif global_outlier:
                status = "Review: Global Ratio Outlier"; bump("global_outlier")
            elif neighborhood_outlier:
                status = "Review: Neighborhood Ratio Outlier"; bump("neighborhood_outlier")
        if not status:
            status = "Verified Compliant"; bump("compliant")
        statuses.append(status)

    out[STATUS_COL] = statuses
    summary = {"checker": "Quality Checker", "total_rows": len(out), **counts}
    summary["issues_requiring_review"] = len(out) - counts.get("compliant", 0)
    summary["valid_ratios"] = ratio_summary.get("valid_ratios", 0)
    return out, summary


def cod(values: Sequence[float]) -> float:
    arr = pd.to_numeric(pd.Series(values), errors="coerce").dropna().to_numpy(dtype=float)
    if len(arr) == 0:
        return 0.0
    median = float(np.median(arr))
    if median <= 0:
        return 0.0
    return float(np.mean(np.abs(arr - median)) / median)


def _stats_row(label: object, grp: pd.DataFrame, price_col: Optional[str], total_col: Optional[str]) -> Dict[str, object]:
    quals = [classify_row_qualification(grp, idx)[0] for idx in grp.index]
    ratio_col = first_col(grp, "ratio")
    ratios = pd.to_numeric(grp[ratio_col], errors="coerce") if ratio_col else pd.Series(dtype=float)
    good_mask = pd.Series([q == "GOOD" for q in quals], index=grp.index)
    good_ratios = ratios[good_mask].dropna() if len(ratios) else pd.Series(dtype=float)
    appraisal = pd.to_numeric(grp[total_col], errors="coerce") if total_col else pd.Series(dtype=float)
    sale = pd.to_numeric(grp[price_col], errors="coerce") if price_col else pd.Series(dtype=float)
    # Weighted mean mirrors VBA: sum appraisal / sum sale for available values.
    weighted = float(appraisal[good_mask].sum() / sale[good_mask].sum()) if total_col and price_col and sale[good_mask].sum() else 0.0
    mean = float(good_ratios.mean()) if len(good_ratios) else 0.0
    return {
        "Label": label,
        "Total Count": len(grp),
        "Good Sales Count": sum(q == "GOOD" for q in quals),
        "Bad Sales Count": sum(q == "BAD" for q in quals),
        "Undetermined Count": sum(q not in {"GOOD", "BAD"} for q in quals),
        "Good Ratio Count": int(good_ratios.count()),
        "Mean": mean,
        "Median": float(good_ratios.median()) if len(good_ratios) else 0.0,
        "Weighted Mean": weighted,
        "Minimum": float(good_ratios.min()) if len(good_ratios) else 0.0,
        "Maximum": float(good_ratios.max()) if len(good_ratios) else 0.0,
        "Range": float(good_ratios.max() - good_ratios.min()) if len(good_ratios) else 0.0,
        "PRD": (mean / weighted) if weighted else 0.0,
        "COD": cod(good_ratios.tolist()),
        "Min Appraisal": float(appraisal[good_mask].min()) if total_col and good_mask.any() else 0.0,
        "Max Appraisal": float(appraisal[good_mask].max()) if total_col and good_mask.any() else 0.0,
        "Average Appraisal": float(appraisal[good_mask].mean()) if total_col and good_mask.any() else 0.0,
        "Average Sale Price": float(sale[good_mask].mean()) if price_col and good_mask.any() else 0.0,
    }


def generate_statistics(df: pd.DataFrame) -> Tuple[pd.DataFrame, Dict[str, pd.DataFrame], Dict[str, object]]:
    # Ensure ratios exist first.
    out = df.copy()
    if first_col(out, "ratio") is None:
        out, _ = ratio_checker(out)
    use_col, nbhd_col, price_col, total_col = first_col(out, "use_code"), first_col(out, "neighborhood"), first_col(out, "sale_price"), first_col(out, "total_value")
    ratio_col = first_col(out, "ratio")

    use_rows = []
    if use_col:
        labels = out[use_col].map(lambda x: normalize_code(x) or "Unknown")
        for label in sorted(labels.unique(), key=lambda x: (not str(x).isdigit(), str(x))):
            grp = out.loc[labels == label]
            r = _stats_row(label, grp, price_col, total_col)
            r["Use Code"] = r.pop("Label")
            use_rows.append(r)
    use_stats = pd.DataFrame(use_rows)

    zone_rows = []
    if nbhd_col:
        labels = out[nbhd_col].map(lambda x: clean_text(x) or "Unknown")
        for label in sorted(labels.unique()):
            grp = out.loc[labels == label]
            r = _stats_row(label, grp, price_col, total_col)
            r["Valuation Zone"] = r.pop("Label")
            zone_rows.append(r)
    zone_stats = pd.DataFrame(zone_rows)

    all_quals = pd.Series([classify_row_qualification(out, idx)[0] for idx in out.index], index=out.index)
    all_ratios = pd.to_numeric(out[ratio_col], errors="coerce") if ratio_col else pd.Series(dtype=float)
    good_ratios = all_ratios[all_quals == "GOOD"].dropna()
    overall_median = float(good_ratios.median()) if len(good_ratios) else 0.0
    overall_cod = cod(good_ratios.tolist())
    certification = "VERIFY PASS" if NATURAL_MIN_RATIO <= overall_median <= NATURAL_MAX_RATIO and overall_cod <= MAX_COD else "REVIEW REQUIRED"
    summary = {
        "checker": "Generate Statistics",
        "total_rows": len(out),
        "good_ratio_count": int(good_ratios.count()),
        "overall_median": overall_median,
        "overall_cod": overall_cod,
        "county_study_certification": certification,
        "zones_with_10_plus_good_sales": int((zone_stats.get("Good Sales Count", pd.Series(dtype=int)) >= MIN_REQUIRED_SALES).sum()) if not zone_stats.empty else 0,
    }
    return out, {"Use Code Statistics": use_stats, "Valuation Zone Statistics": zone_stats}, summary


STATUS_FILL = {
    "critical": "F8A3A3",
    "error": "F8A3A3",
    "review": "FFE3A3",
    "unknown": "D5DDE8",
    "bad sale": "FFF0B5",
    "verified": "D9F2D9",
    "perfect": "D9F2D9",
    "n/a": "E8E8E8",
}


def status_fill(status: object) -> Optional[str]:
    s = clean_text(status).lower()
    for key, color in STATUS_FILL.items():
        if key in s:
            return color
    if "too low" in s or "too high" in s:
        return "E7D5FF"
    return None


def export_results_xlsx(df: pd.DataFrame, summary: Dict[str, object], extra_tables: Optional[Dict[str, pd.DataFrame]] = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Analysis"
    headers = list(df.columns)
    for j, h in enumerate(headers, 1):
        cell = ws.cell(1, j, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="left")
    for i, (_, row) in enumerate(df.iterrows(), 2):
        status = row.get(STATUS_COL, "")
        fill_color = status_fill(status)
        for j, h in enumerate(headers, 1):
            val = row[h]
            if pd.isna(val): val = None
            if isinstance(val, pd.Timestamp): val = val.to_pydatetime()
            cell = ws.cell(i, j, val)
            if fill_color:
                cell.fill = PatternFill("solid", fgColor=fill_color)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for j, h in enumerate(headers, 1):
        max_len = min(35, max(len(str(h)), max((len(str(ws.cell(i, j).value or "")) for i in range(2, min(ws.max_row, 250) + 1)), default=0)) + 2)
        ws.column_dimensions[get_column_letter(j)].width = max_len

    dash = wb.create_sheet("Dashboard")
    dash["A1"] = "Sales Ratio Quality Checker Dashboard"
    dash["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    dash["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    dash.merge_cells("A1:D1")
    dash["A3"] = "Generated"
    dash["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    r = 5
    for k, v in summary.items():
        dash.cell(r, 1, str(k).replace("_", " ").title()).font = Font(bold=True)
        dash.cell(r, 2, v)
        r += 1
    dash.column_dimensions["A"].width = 34
    dash.column_dimensions["B"].width = 50

    for name, table in (extra_tables or {}).items():
        title = name[:31]
        s = wb.create_sheet(title)
        if table is None or table.empty:
            s["A1"] = "No data"
            continue
        for j, h in enumerate(table.columns, 1):
            c = s.cell(1, j, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        for i, (_, row) in enumerate(table.iterrows(), 2):
            for j, h in enumerate(table.columns, 1):
                val = row[h]
                if pd.isna(val): val = None
                s.cell(i, j, val)
        s.freeze_panes = "A2"
        s.auto_filter.ref = s.dimensions
        for j, h in enumerate(table.columns, 1):
            s.column_dimensions[get_column_letter(j)].width = min(30, max(12, len(str(h)) + 2))

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
